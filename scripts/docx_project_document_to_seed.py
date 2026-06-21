from __future__ import annotations

import json
from pathlib import Path

from docx import Document


SOURCE = Path("/Users/yikaihuang/Desktop/2627 AI委員會/2627_AI委員會(創造持衡的影響力)/05_RotaryPassport2.0_Wesly/Rotary_Passport_2.0_功能清單_完整版.docx")
OUT_SQL = Path("supabase/seed/rotary_passport_project_document.sql")
PREFIX = "<!-- rotary-project-document-v1 -->"


def clean(value: str) -> str:
    return " ".join(value.replace("\u3000", " ").split()).strip()


def rows(table):
    return [[clean(cell.text) for cell in row.cells] for row in table.rows]


def dollar_quote(value: str, tag: str) -> str:
    while f"${tag}$" in value:
        tag = f"{tag}_x"
    return f"${tag}${value}${tag}$"


def feature_rows(table_rows: list[list[str]], phase: str):
    features = []
    for row in table_rows[1:]:
        if len(row) < 5 or not row[0]:
            continue
        features.append({
            "no": row[0],
            "name": row[1],
            "description": row[2],
            "permission": row[3],
            "source": row[4],
            "phase": phase,
            "status": "todo",
        })
    return features


def main() -> None:
    doc = Document(SOURCE)
    paragraphs = [clean(p.text) for p in doc.paragraphs]
    paragraphs = [p for p in paragraphs if p]
    tables = [rows(table) for table in doc.tables]

    background_intro = paragraphs[5]
    background_bullets = paragraphs[6:11]
    core_intro = paragraphs[12]
    core_notes = paragraphs[13:15]
    extended_intro = paragraphs[16]
    extended_notes = paragraphs[17:18]
    permission_intro = paragraphs[19]
    data_source_sections = [
        {
            "id": "data-source-vocational-service",
            "title": "5-1 職業服務網介接",
            "bullets": paragraphs[22:25],
        },
        {
            "id": "data-source-public-service",
            "title": "5-2 扶輪公益網連動",
            "bullets": paragraphs[26:28],
        },
        {
            "id": "data-source-club-maintained",
            "title": "5-3 各社自維護資料",
            "bullets": paragraphs[29:32],
        },
        {
            "id": "data-source-district",
            "title": "5-4 地區層級資料",
            "bullets": paragraphs[33:35],
        },
    ]
    privacy_intro = paragraphs[36]
    discussion_intro = paragraphs[38]
    discussion_bullets = paragraphs[39:45]

    project_document = {
        "title": "Rotary Passport 2.0 需求與功能清單彙整",
        "subtitle": "AI委員會工作文件 | 供總監及委員討論使用",
        "source": str(SOURCE),
        "sections": [
            {
                "id": "background",
                "title": "一、專案背景",
                "intro": background_intro,
                "bullets": background_bullets,
            },
            {
                "id": "core-features",
                "title": "二、核心功能清單（14 項）",
                "intro": core_intro,
                "bullets": core_notes,
                "features": feature_rows(tables[0], "核心功能"),
            },
            {
                "id": "extended-features",
                "title": "三、延伸功能清單（5 項）",
                "intro": extended_intro,
                "bullets": extended_notes,
                "features": feature_rows(tables[1], "延伸功能"),
            },
            {
                "id": "permissions",
                "title": "四、權限層級設計",
                "intro": permission_intro,
                "permissions": [
                    {
                        "level": row[0],
                        "role": row[1],
                        "scope": row[2],
                        "notes": row[3],
                    }
                    for row in tables[2][1:]
                ],
            },
            {
                "id": "data-sources",
                "title": "五、資料來源與整合方式",
                "intro": "",
                "bullets": [],
            },
            *data_source_sections,
            {
                "id": "privacy",
                "title": "六、隱私與授權注意事項",
                "intro": privacy_intro,
                "privacy": [
                    {
                        "topic": row[0],
                        "description": row[1],
                        "recommendation": row[2],
                    }
                    for row in tables[3][1:]
                ],
            },
            {
                "id": "discussion",
                "title": "七、待討論事項",
                "intro": discussion_intro,
                "bullets": discussion_bullets,
            },
        ],
    }

    notes = f"{PREFIX}\n{json.dumps(project_document, ensure_ascii=False, indent=2)}"
    background = project_document["sections"][0]
    background_text = "\n".join([background["intro"], *background["bullets"]])

    sql = f"""-- ============================================================
-- Rotary Passport 2.0 功能清單 DOCX 匯入為可編輯專案分段資料
-- 來源：{SOURCE}
-- 執行方式：貼到 Supabase SQL Editor -> Run
-- ============================================================

update public.projects
set
  project_purpose = coalesce(project_purpose, '開發Rotary Passport 2.0，採分階段導入方式，先盤點完整功能、權限、資料來源與隱私授權設計。'),
  project_background = {dollar_quote(notes, "project_document")},
  description = coalesce(description, {dollar_quote(background_text, "project_description")})
where slug = 'rotary-passport-2';
"""
    OUT_SQL.write_text(sql, encoding="utf-8")


if __name__ == "__main__":
    main()
